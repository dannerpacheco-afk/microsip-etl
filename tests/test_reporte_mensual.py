"""Tests for reporte_mensual.py — queries, xlsx output, e-mail."""

from __future__ import annotations

import os
import sys
from datetime import date, datetime
from decimal import Decimal
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from reporte_mensual import (
    CORTES,
    NOTA,
    enviar_por_correo,
    generar_reporte,
    nombre_archivo,
    sql_inventario,
    sql_ventas,
)

MES = date(2026, 8, 1)

VENTAS = [
    {"MES": MES, "CORTE": "Cash&Carry", "CLAVE_PRINCIPAL": "A001", "CODIGO_BARRAS": "7501234567890",
     "DESCRIPCION": "ACEITE 1L", "UNIDAD_VENTA": "PZA", "PROVEEDOR": "PROV SA",
     "PIEZAS": Decimal("10"), "IMPORTE_SIN_IMPUESTOS": Decimal("100.00"),
     "IMPUESTOS": Decimal("16.00"), "IMPORTE_CON_IMPUESTOS": Decimal("116.00")},
    {"MES": MES, "CORTE": "Cash&Carry", "CLAVE_PRINCIPAL": "A002", "CODIGO_BARRAS": None,
     "DESCRIPCION": "AZUCAR 1K", "UNIDAD_VENTA": "PZA", "PROVEEDOR": None,
     "PIEZAS": Decimal("5"), "IMPORTE_SIN_IMPUESTOS": Decimal("50.00"),
     "IMPUESTOS": Decimal("0"), "IMPORTE_CON_IMPUESTOS": Decimal("50.00")},
    {"MES": MES, "CORTE": "Mayoreo tradicional", "CLAVE_PRINCIPAL": "A001", "CODIGO_BARRAS": "7501234567890",
     "DESCRIPCION": "ACEITE 1L", "UNIDAD_VENTA": "PZA", "PROVEEDOR": "PROV SA",
     "PIEZAS": Decimal("-2"), "IMPORTE_SIN_IMPUESTOS": Decimal("-20.00"),
     "IMPUESTOS": Decimal("-3.20"), "IMPORTE_CON_IMPUESTOS": Decimal("-23.20")},
]
INVENTARIO = [
    {"MES": MES, "ALMACEN": "MATRIZ", "CLAVE_PRINCIPAL": "A001", "CODIGO_BARRAS": "7501234567890",
     "DESCRIPCION": "ACEITE 1L", "UNIDAD_VENTA": "PZA", "PROVEEDOR": "PROV SA",
     "EXISTENCIA_FIN_MES": Decimal("120"), "VALOR_COSTO_FIN_MES": Decimal("960.50")},
    {"MES": MES, "ALMACEN": "BODEGA", "CLAVE_PRINCIPAL": "A002", "CODIGO_BARRAS": None,
     "DESCRIPCION": "AZUCAR 1K", "UNIDAD_VENTA": "PZA", "PROVEEDOR": None,
     "EXISTENCIA_FIN_MES": Decimal("30"), "VALOR_COSTO_FIN_MES": Decimal("210")},
]


@pytest.fixture
def loader():
    ld = MagicMock()
    ld.project_id = "proj"
    ld.execute.side_effect = [VENTAS, INVENTARIO]
    return ld


class TestSql:
    def test_ventas_uses_corte_column_and_incluir(self):
        sql = sql_ventas("proj", "ds", "zona", "ALMACENES PACHECO")
        assert "`proj.ds.v_iscam_ventas_mensual`" in sql
        assert "ZONA AS CORTE" in sql
        assert "WHERE MES = @mes AND INCLUIR AND EMPRESA = @empresa" in sql

    def test_ventas_without_empresa_has_no_filter(self):
        sql = sql_ventas("proj", "ds", "formato", "")
        assert "@empresa" not in sql
        assert "FORMATO AS CORTE" in sql

    def test_invalid_corte(self):
        with pytest.raises(ValueError, match="corte inválido"):
            sql_ventas("proj", "ds", "cliente", "")

    def test_inventario_view(self):
        sql = sql_inventario("proj", "ds", "X")
        assert "`proj.ds.v_iscam_inventario_mensual`" in sql
        assert "WHERE MES = @mes" in sql


class TestNombreArchivo:
    def test_strips_spaces(self):
        assert nombre_archivo("ALMACENES PACHECO", date(2026, 8, 15)) == "ISCAM_ALMACENESPACHECO_2026-08.xlsx"

    def test_empty_empresa(self):
        assert nombre_archivo("", MES) == "ISCAM_EMPRESA_2026-08.xlsx"


class TestGenerarReporte:
    def test_writes_workbook_with_three_sheets(self, loader, tmp_path):
        path = generar_reporte(loader, "ds", date(2026, 8, 20), corte="formato",
                               salida=tmp_path / "out", empresa="ALMACENES PACHECO")

        assert path == tmp_path / "out" / "ISCAM_ALMACENESPACHECO_2026-08.xlsx"
        assert path.exists()

        # both queries received the month parameter (first day) and the empresa
        assert loader.execute.call_count == 2
        for c in loader.execute.call_args_list:
            params = c.args[1]
            names = {p.name: p.value for p in params}
            assert names["mes"] == MES
            assert names["empresa"] == "ALMACENES PACHECO"
        assert "FORMATO AS CORTE" in loader.execute.call_args_list[0].args[0]

        wb = load_workbook(path)
        assert wb.sheetnames == ["Ventas", "Inventario", "Resumen"]

        ws = wb["Ventas"]
        headers = [c.value for c in ws[1]]
        assert headers == ["Mes", "Formato", "Clave", "Código de barras", "Descripción", "Unidad",
                           "Proveedor", "Piezas", "Importe sin impuestos", "Impuestos",
                           "Importe con impuestos"]
        assert ws.max_row == 1 + len(VENTAS)
        assert ws["A2"].value == "2026-08"
        assert ws["B2"].value == "Cash&Carry"
        assert ws["D2"].value == "7501234567890"
        assert ws["K2"].value == pytest.approx(116.00)  # openpyxl reads numbers back as float
        assert ws["K2"].number_format == "#,##0.00"
        assert ws["A1"].font.bold
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref == f"A1:K{1 + len(VENTAS)}"
        assert ws.column_dimensions["E"].width > 20

        wi = wb["Inventario"]
        assert [c.value for c in wi[1]] == ["Mes", "Almacén", "Clave", "Código de barras", "Descripción",
                                            "Unidad", "Proveedor", "Existencia fin de mes", "Valor a costo"]
        assert wi.max_row == 1 + len(INVENTARIO)
        assert wi["H2"].value == pytest.approx(120)
        assert wi.freeze_panes == "A2"

        wr = wb["Resumen"]
        values = [[c.value for c in row] for row in wr.iter_rows()]
        flat = [v for row in values for v in row if v is not None]
        assert "Reporte ISCAM" in flat and "ALMACENES PACHECO" in flat
        assert "2026-08" in flat
        assert NOTA in flat
        # totals by corte + grand total
        assert wr["A5"].value == "Formato"
        rows_by_corte = {wr.cell(row=r, column=1).value: r for r in range(6, 9)}
        assert set(rows_by_corte) == {"Cash&Carry", "Mayoreo tradicional", "Total"}
        r = rows_by_corte["Cash&Carry"]
        assert wr.cell(row=r, column=3).value == pytest.approx(15)
        assert wr.cell(row=r, column=6).value == pytest.approx(166.00)
        t = rows_by_corte["Total"]
        assert wr.cell(row=t, column=3).value == pytest.approx(13)
        assert wr.cell(row=t, column=6).value == pytest.approx(142.80)
        assert "Generado" in flat
        gen_row = next(r for r in range(1, wr.max_row + 1) if wr.cell(row=r, column=1).value == "Generado")
        assert isinstance(wr.cell(row=gen_row, column=2).value, datetime)

    def test_corte_zona_changes_header(self, loader, tmp_path):
        path = generar_reporte(loader, "ds", MES, corte="zona", salida=tmp_path, empresa="X")
        ws = load_workbook(path)["Ventas"]
        assert ws["B1"].value == "Zona"
        assert "ZONA AS CORTE" in loader.execute.call_args_list[0].args[0]

    def test_empty_month_still_writes_file(self, tmp_path):
        ld = MagicMock()
        ld.project_id = "proj"
        ld.execute.side_effect = [[], []]
        path = generar_reporte(loader=ld, dataset="ds", mes=MES, salida=tmp_path, empresa="X")
        wb = load_workbook(path)
        assert wb["Ventas"].max_row == 1
        assert wb["Inventario"].max_row == 1

    def test_rejects_bad_corte_before_querying(self, loader, tmp_path):
        with pytest.raises(ValueError):
            generar_reporte(loader, "ds", MES, corte="cliente", salida=tmp_path)
        loader.execute.assert_not_called()

    def test_accepts_bigquery_like_rows(self, tmp_path):
        class Row:  # minimal stand-in for google.cloud.bigquery.table.Row
            def __init__(self, d):
                self._d = d

            def keys(self):
                return self._d.keys()

            def __getitem__(self, k):
                return self._d[k]

        ld = MagicMock()
        ld.project_id = "proj"
        ld.execute.side_effect = [[Row(VENTAS[0])], [Row(INVENTARIO[0])]]
        path = generar_reporte(ld, "ds", MES, salida=tmp_path, empresa="X")
        wb = load_workbook(path)
        assert wb["Ventas"]["E2"].value == "ACEITE 1L"

    def test_all_cortes_supported(self):
        assert set(CORTES) == {"formato", "zona", "almacen"}


class TestEnviarPorCorreo:
    def test_sends_with_starttls_and_attachment(self, tmp_path, caplog):
        path = tmp_path / "ISCAM_X_2026-08.xlsx"
        path.write_bytes(b"PK\x03\x04fake")
        with patch("reporte_mensual.smtplib.SMTP") as smtp_cls:
            smtp = smtp_cls.return_value.__enter__.return_value
            enviar_por_correo(path, ["a@x.mx", "b@x.mx"], host="mail.x.mx", port=587,
                              user="etl", password="s3cret", sender="etl@x.mx")

        smtp_cls.assert_called_once_with("mail.x.mx", 587, timeout=60.0)
        smtp.starttls.assert_called_once()
        smtp.login.assert_called_once_with("etl", "s3cret")
        msg = smtp.send_message.call_args.args[0]
        assert msg["To"] == "a@x.mx, b@x.mx"
        assert msg["From"] == "etl@x.mx"
        assert "ISCAM" in msg["Subject"]
        attachments = [p for p in msg.iter_attachments()]
        assert len(attachments) == 1
        assert attachments[0].get_filename() == "ISCAM_X_2026-08.xlsx"
        assert attachments[0].get_content_type() == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert "s3cret" not in caplog.text

    def test_no_login_without_user(self, tmp_path):
        path = tmp_path / "r.xlsx"
        path.write_bytes(b"x")
        with patch("reporte_mensual.smtplib.SMTP") as smtp_cls:
            smtp = smtp_cls.return_value.__enter__.return_value
            enviar_por_correo(path, ["a@x.mx"], host="h", sender="s@x.mx")
        smtp.login.assert_not_called()

    def test_requires_recipients(self, tmp_path):
        with pytest.raises(ValueError):
            enviar_por_correo(tmp_path / "r.xlsx", [], host="h", sender="s")
