"""Reporte mensual ISCAM: ventas e inventario de un mes en un .xlsx.

    python main.py reporte-mensual [--mes YYYY-MM] [--corte formato|zona|almacen]
                                   [--salida DIR] [--sin-correo]

Fuentes (BigQuery): ``v_iscam_ventas_mensual`` (ventas netas F − D con
impuestos, filtradas por ``INCLUIR``) y ``v_iscam_inventario_mensual``
(existencia y valor a costo al cierre de mes). El archivo se puede enviar por
correo con ``enviar_por_correo`` cuando SMTP está configurado.
"""

from __future__ import annotations

import logging
import re
import smtplib
from collections.abc import Iterable
from datetime import date, datetime
from decimal import Decimal
from email.message import EmailMessage
from pathlib import Path

from google.cloud import bigquery
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)

# corte → (columna en v_iscam_ventas_mensual, encabezado en el Excel)
CORTES: dict[str, tuple[str, str]] = {
    "formato": ("FORMATO", "Formato"),
    "zona": ("ZONA", "Zona"),
    "almacen": ("ALMACEN", "Almacén"),
}
CORTE_DEFAULT = "formato"

NOTA = "Ventas netas de devoluciones; importes en MXN"
_FMT_NUM = "#,##0.00"
_FMT_INT = "#,##0"
_XLSX_MIME = ("application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet")

VENTAS_HEADERS = [
    "Mes",
    None,  # corte, resolved at runtime
    "Clave",
    "Código de barras",
    "Descripción",
    "Unidad",
    "Proveedor",
    "Piezas",
    "Importe sin impuestos",
    "Impuestos",
    "Importe con impuestos",
]
VENTAS_KEYS = [
    "MES", "CORTE", "CLAVE_PRINCIPAL", "CODIGO_BARRAS", "DESCRIPCION", "UNIDAD_VENTA",
    "PROVEEDOR", "PIEZAS", "IMPORTE_SIN_IMPUESTOS", "IMPUESTOS", "IMPORTE_CON_IMPUESTOS",
]
VENTAS_WIDTHS = [10, 26, 14, 18, 48, 10, 32, 12, 20, 14, 22]
VENTAS_NUMERIC = {"PIEZAS": _FMT_NUM, "IMPORTE_SIN_IMPUESTOS": _FMT_NUM,
                  "IMPUESTOS": _FMT_NUM, "IMPORTE_CON_IMPUESTOS": _FMT_NUM}

INVENTARIO_HEADERS = [
    "Mes",
    "Almacén",
    "Clave",
    "Código de barras",
    "Descripción",
    "Unidad",
    "Proveedor",
    "Existencia fin de mes",
    "Valor a costo",
]
INVENTARIO_KEYS = [
    "MES", "ALMACEN", "CLAVE_PRINCIPAL", "CODIGO_BARRAS", "DESCRIPCION", "UNIDAD_VENTA",
    "PROVEEDOR", "EXISTENCIA_FIN_MES", "VALOR_COSTO_FIN_MES",
]
INVENTARIO_WIDTHS = [10, 26, 14, 18, 48, 10, 32, 20, 18]
INVENTARIO_NUMERIC = {"EXISTENCIA_FIN_MES": _FMT_NUM, "VALOR_COSTO_FIN_MES": _FMT_NUM}


# --- Queries ---


def _corte_columna(corte: str) -> tuple[str, str]:
    try:
        return CORTES[corte]
    except KeyError:
        raise ValueError(f"corte inválido {corte!r}; usa uno de {', '.join(CORTES)}") from None


def sql_ventas(project: str, dataset: str, corte: str, empresa: str) -> str:
    columna, _ = _corte_columna(corte)
    filtro_empresa = "AND EMPRESA = @empresa" if empresa else ""
    return f"""
SELECT
  MES,
  {columna} AS CORTE,
  CLAVE_PRINCIPAL,
  CODIGO_BARRAS,
  DESCRIPCION,
  UNIDAD_VENTA,
  PROVEEDOR,
  SUM(PIEZAS) AS PIEZAS,
  SUM(IMPORTE_SIN_IMPUESTOS) AS IMPORTE_SIN_IMPUESTOS,
  SUM(IMPUESTOS) AS IMPUESTOS,
  SUM(IMPORTE_CON_IMPUESTOS) AS IMPORTE_CON_IMPUESTOS
FROM `{project}.{dataset}.v_iscam_ventas_mensual`
WHERE MES = @mes AND INCLUIR {filtro_empresa}
GROUP BY 1, 2, 3, 4, 5, 6, 7
ORDER BY CORTE, DESCRIPCION, CLAVE_PRINCIPAL
"""


def sql_inventario(project: str, dataset: str, empresa: str) -> str:
    filtro_empresa = "AND EMPRESA = @empresa" if empresa else ""
    return f"""
SELECT
  MES,
  ALMACEN,
  CLAVE_PRINCIPAL,
  CODIGO_BARRAS,
  DESCRIPCION,
  UNIDAD_VENTA,
  PROVEEDOR,
  EXISTENCIA_FIN_MES,
  VALOR_COSTO_FIN_MES
FROM `{project}.{dataset}.v_iscam_inventario_mensual`
WHERE MES = @mes
  AND (COALESCE(EXISTENCIA_FIN_MES, 0) <> 0 OR COALESCE(VALOR_COSTO_FIN_MES, 0) <> 0)
  {filtro_empresa}
ORDER BY ALMACEN, DESCRIPCION, CLAVE_PRINCIPAL
"""


def _params(mes: date, empresa: str) -> list:
    params = [bigquery.ScalarQueryParameter("mes", "DATE", mes)]
    if empresa:
        params.append(bigquery.ScalarQueryParameter("empresa", "STRING", empresa))
    return params


def _to_dicts(result: Iterable) -> list[dict]:
    """BigQuery ``Row`` objects (or plain dicts in tests) → list of dicts."""
    return [row if isinstance(row, dict) else dict(row) for row in result]


def consultar_ventas(loader, dataset: str, mes: date, corte: str, empresa: str) -> list[dict]:
    sql = sql_ventas(loader.project_id, dataset, corte, empresa)
    return _to_dicts(loader.execute(sql, _params(mes, empresa)))


def consultar_inventario(loader, dataset: str, mes: date, empresa: str) -> list[dict]:
    sql = sql_inventario(loader.project_id, dataset, empresa)
    return _to_dicts(loader.execute(sql, _params(mes, empresa)))


# --- Excel ---


def nombre_archivo(empresa: str, mes: date) -> str:
    """``ISCAM_<EMPRESA sin espacios>_<YYYY-MM>.xlsx``."""
    limpio = re.sub(r"[^A-Za-z0-9_-]", "", empresa.replace(" ", "")) or "EMPRESA"
    return f"ISCAM_{limpio}_{mes:%Y-%m}.xlsx"


def _mes_texto(value, fallback: date) -> str:
    if isinstance(value, (date, datetime)):
        return f"{value:%Y-%m}"
    if isinstance(value, str) and len(value) >= 7:
        return value[:7]
    return f"{fallback:%Y-%m}"


def _num(value) -> Decimal:
    if value is None:
        return Decimal(0)
    if isinstance(value, Decimal):
        return value
    return Decimal(str(value))


def _cell_value(value):
    """Normalize a BigQuery value for openpyxl."""
    if isinstance(value, Decimal):
        return value if value.is_finite() else None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    return value


def _escribir_hoja(
    ws: Worksheet,
    headers: list[str],
    keys: list[str],
    rows: list[dict],
    widths: list[int],
    numeric: dict[str, str],
    mes: date,
) -> None:
    bold = Font(bold=True)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = bold
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        values = []
        for key in keys:
            value = row.get(key)
            if key == "MES":
                value = _mes_texto(value, mes)
            values.append(_cell_value(value))
        ws.append(values)
    if rows:
        for idx, key in enumerate(keys, start=1):
            fmt = numeric.get(key)
            if fmt:
                for col in ws.iter_cols(min_col=idx, max_col=idx, min_row=2, max_row=ws.max_row):
                    for c in col:
                        c.number_format = fmt
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 1)}"


def _escribir_resumen(
    ws: Worksheet,
    ventas: list[dict],
    inventario: list[dict],
    corte_label: str,
    mes: date,
    empresa: str,
    generado: datetime,
) -> None:
    bold = Font(bold=True)
    title = Font(bold=True, size=13)

    ws["A1"] = "Reporte ISCAM"
    ws["A1"].font = title
    ws["B1"] = empresa
    ws["A2"] = "Mes"
    ws["A2"].font = bold
    ws["B2"] = f"{mes:%Y-%m}"
    ws["A3"] = "Corte"
    ws["A3"].font = bold
    ws["B3"] = corte_label

    # Totales de ventas por corte
    totales: dict[str, list[Decimal]] = {}
    orden: list[str] = []
    for row in ventas:
        key = row.get("CORTE") or "(sin dato)"
        if key not in totales:
            totales[key] = [Decimal(0)] * 4
            orden.append(key)
        acc = totales[key]
        acc[0] += _num(row.get("PIEZAS"))
        acc[1] += _num(row.get("IMPORTE_SIN_IMPUESTOS"))
        acc[2] += _num(row.get("IMPUESTOS"))
        acc[3] += _num(row.get("IMPORTE_CON_IMPUESTOS"))

    header_row = 5
    headers = [corte_label, "Artículos", "Piezas", "Importe sin impuestos", "Impuestos", "Importe con impuestos"]
    for col, text in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=text)
        c.font = bold
    conteo: dict[str, int] = {}
    for row in ventas:
        key = row.get("CORTE") or "(sin dato)"
        conteo[key] = conteo.get(key, 0) + 1
    r = header_row + 1
    grand = [Decimal(0)] * 4
    for key in orden:
        acc = totales[key]
        ws.cell(row=r, column=1, value=key)
        ws.cell(row=r, column=2, value=conteo.get(key, 0)).number_format = _FMT_INT
        for i, val in enumerate(acc, start=3):
            ws.cell(row=r, column=i, value=val).number_format = _FMT_NUM
            grand[i - 3] += val
        r += 1
    ws.cell(row=r, column=1, value="Total").font = bold
    ws.cell(row=r, column=2, value=len(ventas)).number_format = _FMT_INT
    ws.cell(row=r, column=2).font = bold
    for i, val in enumerate(grand, start=3):
        c = ws.cell(row=r, column=i, value=val)
        c.number_format = _FMT_NUM
        c.font = bold
    r += 2

    # Inventario
    ws.cell(row=r, column=1, value="Inventario fin de mes").font = bold
    r += 1
    ws.cell(row=r, column=1, value="Renglones")
    ws.cell(row=r, column=2, value=len(inventario)).number_format = _FMT_INT
    r += 1
    ws.cell(row=r, column=1, value="Existencia total")
    ws.cell(row=r, column=2, value=sum((_num(x.get("EXISTENCIA_FIN_MES")) for x in inventario), Decimal(0))).number_format = _FMT_NUM
    r += 1
    ws.cell(row=r, column=1, value="Valor a costo")
    ws.cell(row=r, column=2, value=sum((_num(x.get("VALOR_COSTO_FIN_MES")) for x in inventario), Decimal(0))).number_format = _FMT_NUM
    r += 2

    ws.cell(row=r, column=1, value="Generado")
    ws.cell(row=r, column=2, value=generado.replace(microsecond=0, tzinfo=None)).number_format = "yyyy-mm-dd hh:mm:ss"
    r += 1
    ws.cell(row=r, column=1, value="Nota")
    ws.cell(row=r, column=2, value=NOTA)

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 26
    for col in "CDEF":
        ws.column_dimensions[col].width = 20


def escribir_xlsx(
    path: Path,
    ventas: list[dict],
    inventario: list[dict],
    corte: str,
    mes: date,
    empresa: str,
    generado: datetime | None = None,
) -> Path:
    """Write the three sheets (Ventas, Inventario, Resumen) to ``path``."""
    _, corte_label = _corte_columna(corte)
    generado = generado or datetime.now()
    wb = Workbook()
    ws_v = wb.active
    ws_v.title = "Ventas"
    headers_v = [h if h is not None else corte_label for h in VENTAS_HEADERS]
    _escribir_hoja(ws_v, headers_v, VENTAS_KEYS, ventas, VENTAS_WIDTHS, VENTAS_NUMERIC, mes)

    ws_i = wb.create_sheet("Inventario")
    _escribir_hoja(ws_i, INVENTARIO_HEADERS, INVENTARIO_KEYS, inventario, INVENTARIO_WIDTHS, INVENTARIO_NUMERIC, mes)

    ws_r = wb.create_sheet("Resumen")
    _escribir_resumen(ws_r, ventas, inventario, corte_label, mes, empresa, generado)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def generar_reporte(
    loader,
    dataset: str,
    mes: date,
    corte: str = CORTE_DEFAULT,
    salida: Path | str = Path("reports"),
    empresa: str = "",
) -> Path:
    """Query both ISCAM views for ``mes`` and write the .xlsx. Returns its path."""
    mes = mes.replace(day=1)
    _corte_columna(corte)  # validate early
    logger.info("Reporte ISCAM %s corte=%s empresa=%s", f"{mes:%Y-%m}", corte, empresa or "(todas)")
    ventas = consultar_ventas(loader, dataset, mes, corte, empresa)
    logger.info("  ventas: %d renglones", len(ventas))
    inventario = consultar_inventario(loader, dataset, mes, empresa)
    logger.info("  inventario: %d renglones", len(inventario))
    if not ventas:
        logger.warning("Sin ventas para %s; el reporte saldrá vacío", f"{mes:%Y-%m}")
    path = Path(salida) / nombre_archivo(empresa, mes)
    escribir_xlsx(path, ventas, inventario, corte, mes, empresa)
    logger.info("Reporte escrito en %s", path)
    return path


# --- Correo ---


def enviar_por_correo(
    path: Path | str,
    to: list[str],
    host: str,
    port: int = 587,
    user: str = "",
    password: str = "",
    sender: str = "",
    subject: str | None = None,
    body: str | None = None,
    timeout: float = 60.0,
) -> None:
    """Send ``path`` as an attachment via SMTP + STARTTLS.

    The password is never logged. Raises ``smtplib.SMTPException`` / ``OSError``
    on failure so the caller can decide the exit code.
    """
    path = Path(path)
    if not to:
        raise ValueError("no recipients")
    if not host or not sender:
        raise ValueError("smtp host and sender are required")

    msg = EmailMessage()
    msg["From"] = sender
    msg["To"] = ", ".join(to)
    msg["Subject"] = subject or f"Reporte ISCAM {path.stem.replace('ISCAM_', '').replace('_', ' ')}"
    msg.set_content(body or f"Adjunto {path.name}.\n\n{NOTA}.\n")
    maintype, subtype = _XLSX_MIME
    msg.add_attachment(path.read_bytes(), maintype=maintype, subtype=subtype, filename=path.name)

    logger.info("Enviando %s a %s vía %s:%d (user=%s)", path.name, to, host, port, user or "-")
    with smtplib.SMTP(host, port, timeout=timeout) as smtp:
        smtp.ehlo()
        smtp.starttls()
        smtp.ehlo()
        if user:
            smtp.login(user, password)
        smtp.send_message(msg)
    logger.info("Correo enviado a %s", ", ".join(to))
